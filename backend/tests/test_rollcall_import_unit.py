"""Roll call import (core/rollcall_import.py): week tabs and their dates, the
sheet's colours, first-name matching, and "the app wins" — only untouched
Bells days are filled.

Pure in-memory: a small workbook built here, mongomock_motor. Placeholder names.
"""
import asyncio
import io
import os
import sys
from datetime import date
from pathlib import Path

os.environ.setdefault("MONGO_URL", "mongodb://127.0.0.1:1/unused")
BACKEND_DIR = Path(__file__).resolve().parents[1]
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

import openpyxl  # noqa: E402
from openpyxl.styles import PatternFill  # noqa: E402
from bson import ObjectId  # noqa: E402
from mongomock_motor import AsyncMongoMockClient  # noqa: E402

from core import rollcall_import as ri  # noqa: E402

run = asyncio.run
GOLD, SILVER = "FFFFD966", "FFCCCCCC"
BLACK, GREY, YELLOW, RED = "FF000000", "FF666666", "FFFFFF00", "FFFF0000"


def _fill(c, rgb):
    c.fill = PatternFill(fill_type="solid", fgColor=rgb)


def _book() -> bytes:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("WE 2792026")
    header = ["", "NAME", "ST", "PR", "Breakeven", "Last Week", "Goal", "Focus", "Was", "Goal", "Team Goal"]
    for i, h in enumerate(header, start=1):
        ws.cell(1, i, h)
    for di, dn in enumerate(["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"]):
        ws.cell(1, 12 + di * 2, dn)
    ws.cell(1, 26, "TOTAL GOLD")

    def person(r, team, name, days):
        if team:
            ws.cell(r, 1, team)
        ws.cell(r, 2, name)
        for di, d in enumerate(days):
            g, s = ws.cell(r, 12 + di * 2), ws.cell(r, 13 + di * 2)
            if isinstance(d, tuple):
                g.value, s.value = d
                _fill(g, GOLD); _fill(s, SILVER)
            elif d:
                _fill(g, d); _fill(s, d)

    person(3, "Team Example", "Example A", [(3, 1), (0, 0), BLACK, GREY, YELLOW, RED, GREY])
    person(4, "", "Example B", [None, (2, 0), None, None, None, None, (1, 0)])
    person(5, "", "Nobody Here", [(1, 0)] + [None] * 6)
    ws.cell(6, 2, "Total")
    person(7, "", "After Total", [(9, 9)] + [None] * 6)
    wb.create_sheet("WE 9_10 - Datum _ Fundraising")       # another campaign: ignored
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def test_week_ending_from_tab_names():
    assert ri.week_ending_of("WE 2792026") == date(2026, 9, 27)
    assert ri.week_ending_of("WE 692026") == date(2026, 9, 6)
    assert ri.week_ending_of("WE 11102026") == date(2026, 10, 11)
    assert ri.week_ending_of("WE 04102026") == date(2026, 10, 4)
    assert ri.week_ending_of("WE 9_10 - Datum _ Fundraising") is None
    assert ri.week_ending_of("WE 1102026") is None          # 1 Oct 2026 isn't a Sunday


def test_colours():
    assert ri.status_from_colour((0, 0, 0)) == "ns"
    assert ri.status_from_colour((0x66, 0x66, 0x66)) == "ab"
    assert ri.status_from_colour((0xFF, 0xFF, 0x00)) == "pc"
    assert ri.status_from_colour((0xFF, 0x00, 0x00)) == "skip"
    assert ri.status_from_colour((0xFF, 0xD9, 0x66)) is None   # gold cell background
    assert ri.status_from_colour((0xCC, 0xCC, 0xCC)) is None   # silver cell background


def test_reads_a_week_tab():
    weeks = ri.read_workbook(_book(), date(2026, 9, 1), date(2026, 10, 11))
    assert list(weeks) == ["2026-09-27"]
    rows = {r["name"]: r for r in weeks["2026-09-27"]}
    assert set(rows) == {"Example A", "Example B", "Nobody Here"}   # stops at Total
    a = rows["Example A"]["days"]
    assert a[0] == {"over30": 3, "under30": 1, "memberships": None, "status": "in"}
    assert a[1]["status"] == "in" and a[1]["over30"] == 0          # a typed 0 is a day in
    assert [a[2]["status"], a[3]["status"], a[4]["status"]] == ["ns", "ab", "pc"]
    assert a[5] is None                                            # red: left alone
    assert a[6] is None                                            # grey Sunday: nothing
    assert rows["Example B"]["days"][6]["over30"] == 1             # a Sunday number counts
    assert rows["Example A"]["team"] == "Team Example"


def test_match_by_first_name():
    people = [{"id": "1", "name": "Example Alpha"}, {"id": "2", "name": "Example Beta"},
              {"id": "3", "name": "Sam Person"}, {"id": "4", "name": "Sam Other"}]
    assert ri.match_name("Example Alpha", people)[0]["id"] == "1"
    assert ri.match_name("Sam P", people)[0]["id"] == "3"           # surname initial narrows it
    assert ri.match_name("Sam", people)[0] is None                  # two Sams: skipped
    assert ri.match_name("Zed", people)[0] is None
    # A short name in capitals: initials, or the start of a first name.
    shorts = [{"id": "5", "name": "Abcde Placeholder"}, {"id": "6", "name": "Quinn Rivers"}]
    assert ri.match_name("AB", shorts)[0]["id"] == "5"
    assert ri.match_name("QR", shorts)[0]["id"] == "6"
    assert ri.match_name("ZZ", shorts)[0] is None
    assert ri.match_name("Ab", shorts)[0] is None           # only capitals count as a short name


def test_merge_keeps_the_app():
    app = [{"over30": None, "under30": 2, "memberships": None, "status": "in"},   # typed in the app
           {"over30": None, "under30": None, "memberships": None, "status": "ab"}, # set in the app
           {"over30": None, "under30": None, "memberships": None, "status": "off"}]
    sheet = [{"over30": 5, "under30": 0, "memberships": None, "status": "in"},
             {"over30": 1, "under30": 0, "memberships": None, "status": "in"},
             {"over30": None, "under30": None, "memberships": None, "status": "ns"}] + [None] * 4
    days, filled = ri.merge_days(app, sheet)
    assert filled == 1
    assert days[0]["under30"] == 2 and days[1]["status"] == "ab" and days[2]["status"] == "ns"


def _db():
    db = AsyncMongoMockClient()["t"]
    a, b = ObjectId(), ObjectId()
    run(db.users.insert_many([
        {"_id": a, "name": "Example A", "office_id": "o1", "role": "trainee"},
        {"_id": b, "name": "Example B", "office_id": "o1", "role": "trainee"},
    ]))
    # B already has the week, with Tuesday typed in the app.
    days = [{"over30": None, "under30": None, "memberships": None, "status": "off"} for _ in range(7)]
    days[1] = {"over30": None, "under30": 4, "memberships": None, "status": "in"}
    run(db.bells_entries.insert_one({"id": "b-week", "office_id": "o1", "week_ending": "2026-09-27",
                                     "user_id": str(b), "user_name": "Example B", "days": days}))
    return db, str(a), str(b)


def _totals(entry, fees):
    return {"earnings": float(sum((d.get("over30") or 0) + (d.get("under30") or 0) for d in entry["days"]))}


def test_preview_then_import():
    db, a, b = _db()
    weeks = ri.read_workbook(_book(), date(2026, 9, 1), date(2026, 10, 11))
    plan = run(ri.plan_import(db, "o1", weeks))
    assert [u["sheet_name"] for u in plan["unmatched"]] == ["Nobody Here"]
    wk = plan["weeks"][0]
    assert wk["days"] == 6 and wk["new_rows"] == 1                 # A: 5 days new row; B: Sunday only
    assert run(db.bells_entries.count_documents({})) == 1          # preview wrote nothing

    out = run(ri.apply_import(db, "o1", weeks, None, _totals, "admin-1"))
    assert out == {"days_filled": 6, "rows_updated": 1, "rows_created": 1}
    ea = run(db.bells_entries.find_one({"user_id": a}))
    assert ea["created_by_id"] == "rollcall-import"
    assert [d["status"] for d in ea["days"]] == ["in", "in", "ns", "ab", "pc", "off", "off"]
    assert ea["earnings"] == 4.0
    eb = run(db.bells_entries.find_one({"user_id": b}))
    assert eb["days"][1]["under30"] == 4                           # the app's Tuesday kept
    assert eb["days"][6]["over30"] == 1

    again = run(ri.apply_import(db, "o1", weeks, None, _totals, "admin-1"))
    assert again["days_filled"] == 0                               # nothing left to fill
