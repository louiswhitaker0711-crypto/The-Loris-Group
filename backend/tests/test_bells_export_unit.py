"""Bells as Excel (core/bells_export.py): a tab per week, teams with totals,
days as the app shows them, EX from the exit day; Admins only."""
import asyncio
import io
import os
import sys
from pathlib import Path

os.environ.setdefault("MONGO_URL", "mongodb://127.0.0.1:1/unused")
BACKEND_DIR = Path(__file__).resolve().parents[1]
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

import openpyxl  # noqa: E402
import pytest  # noqa: E402
from fastapi import HTTPException  # noqa: E402

from core.bells_export import build_workbook  # noqa: E402

WEEK = {"week_ending": "2026-10-25", "mini_team_goal": 40, "entries": [
    {"user_id": "u1", "user_name": "Example Person", "recruit_badge": "PR", "weekly_goal": 12, "last_week_total": 16,
     "days": [{"status": "in", "over30": 3, "under30": 1}, {"status": "ab"}, {"status": "ns"}, {"status": "off"},
              {"status": "off"}, {"status": "pc"}, {"status": "ab"}],
     "total_sales": 4, "total_over30": 3, "total_under30": 1, "days_worked": 1, "earnings": 80.0, "mc_fees": 76.0},
    {"user_id": "u2", "user_name": "Sample Person", "left_on": "2026-10-21",
     "primary_team": {"leader_id": "u3", "team_name": "Example Team"},
     "days": [{"status": "in", "over30": 2}] + [{"status": "off"}] * 6, "total_sales": 2, "days_worked": 1, "earnings": 40.0},
]}


def _rows(data, tab=0):
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb[wb.sheetnames[tab]]
    return wb.sheetnames, [list(r) for r in ws.iter_rows(values_only=True)]


def test_a_tab_per_week_like_the_sheet():
    names, rows = _rows(build_workbook([WEEK, {"week_ending": "2026-10-18", "entries": []}], "Example Office", "The Example Group"))
    assert names == ["WE 25 Oct 26", "WE 18 Oct 26"]
    assert rows[2][:8] == ["Name", "PR/OR", "Coach", "Stage", "Goal", "LW", "B/E", "Mon"] and rows[2][-1] == "MC fees"
    person = next(r for r in rows if r[0] == "Example Person")
    assert person[1] == "PR" and person[7:14] == [4, "NS", "X", None, None, "IO", "NS"]
    gone = next(r for r in rows if r[0] == "Sample Person")
    assert gone[7:14] == [2, None, "EX", "EX", "EX", "EX", "EX"]                      # exited Wed 21 Oct
    assert rows[-1][0] == "Example Office total" and rows[-1][14] == 6


def test_no_weeks_still_makes_a_file():
    names, rows = _rows(build_workbook([]))
    assert names == ["Bells"] and rows[0][0] == "No Bells weeks yet."


def test_only_admins_export(monkeypatch):
    from routes import bells

    async def coach(_r):
        return {"id": "c1", "role": "leader", "office_id": "o1"}
    monkeypatch.setattr(bells, "get_current_user", coach)
    with pytest.raises(HTTPException) as err:
        asyncio.run(bells.export_bells(object()))
    assert err.value.status_code == 403
