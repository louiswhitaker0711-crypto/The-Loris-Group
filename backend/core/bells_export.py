"""Bells as an Excel workbook: one tab per week, newest first, laid out like
the sheet in the app (routes/bells.py → GET /api/bells/export, Admins only).

Each tab: the office and week, then one block per team (its people in the
sheet's order, then the team's total), then the office total. Days read the
way the app shows them: sign-ups on a day in (coloured 0-1 red, 2 amber,
3 light green, 4+ dark green), NS / X / IO / RT / NC, EX from the day someone
exited, blank for a day off. Goal and LW are tinted by the weekly zone
(core/week_bands). Fees as the Admin sees them, MC fees included.

Pure: takes the week payloads (routes.bells.bells_week_for) and returns bytes.
"""
from __future__ import annotations

import io
from datetime import date, timedelta

from core.week_bands import week_band

DAYS = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]
STATUS = {"ab": "NS", "ns": "X", "pc": "IO", "rt": "RT", "nc": "NC", "sd": "SD"}

# Fill / text colours, matching the app's light theme.
DAY_LOOK = {"red": ("FDE2E2", "B91C1C"), "amber": ("FCEBD0", "B45309"),
            "light_green": ("D3F5DE", "15803D"), "dark_green": ("166534", "FFFFFF")}
WEEK_LOOK = {"red": ("FDE2E2", "B91C1C"), "amber": ("FCEBD0", "B45309"),
             "green": ("D3F5DE", "15803D"), "super_green": ("166534", "FFFFFF")}
STATUS_LOOK = {"NS": ("E5E7EB", "374151"), "X": ("111111", "FFFFFF"), "IO": (None, "0E7490"),
               "RT": (None, "B45309"), "NC": (None, "6B7280"), "EX": ("DC2626", "FFFFFF")}
HEAD_FILL = "1E3A8A"
TEAM_FILL = "E8EEF8"
TOTAL_FILL = "DBE6F0"


def day_band(n: int) -> str:
    return "dark_green" if n >= 4 else "light_green" if n == 3 else "amber" if n == 2 else "red"


def tab_name(week_ending: str, used: set) -> str:
    try:
        d = date.fromisoformat(week_ending)
        name = f"WE {d.day} {d.strftime('%b')} {d.strftime('%y')}"
    except ValueError:
        name = (week_ending or "Week")[:28]
    base, n = name, 2
    while name in used:
        name = f"{base} ({n})"
        n += 1
    used.add(name)
    return name[:31]


def uk_date(iso: str) -> str:
    try:
        d = date.fromisoformat(iso)
        return f"{d.day} {d.strftime('%b %Y')}"
    except ValueError:
        return iso or ""


def _num(v) -> float:
    try:
        return float(v or 0)
    except (TypeError, ValueError):
        return 0.0


def _exited_on(e: dict, week_ending: str, i: int) -> bool:
    left = str(e.get("left_on") or "")[:10]
    if not left:
        return False
    try:
        day = date.fromisoformat(week_ending) - timedelta(days=6 - i)
    except ValueError:
        return False
    return day.isoformat() >= left


def day_cell(e: dict, d: dict, week_ending: str, i: int):
    """(value, look) for one day: a number for a day in, a status code, or blank."""
    if _exited_on(e, week_ending, i):
        return "EX", STATUS_LOOK["EX"]
    st = (d or {}).get("status") or "off"
    if st == "in":
        n = int(_num(d.get("over30")) + _num(d.get("under30")))
        return n, DAY_LOOK[day_band(n)]
    if st in STATUS:
        code = STATUS[st]
        return code, STATUS_LOOK[code]
    return "", None


def team_of(e: dict, org_name: str) -> str:
    pt = e.get("primary_team") or {}
    return (pt.get("team_name") or "").strip() or org_name


def build_workbook(weeks: list[dict], office_name: str = "", org_name: str = "The Office") -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    used: set = set()
    thin = Side(style="thin", color="CBD8E6")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    centre = Alignment(horizontal="center", vertical="center")

    def paint(cell, look, bold=False):
        fill, colour = look if look else (None, None)
        if fill:
            cell.fill = PatternFill("solid", fgColor=fill)
        cell.font = Font(bold=bold, color=colour or "111827")

    if not weeks:
        ws = wb.create_sheet("Bells")
        ws["A1"] = "No Bells weeks yet."
    for wk in weeks:
        we = wk.get("week_ending") or ""
        entries = [e for e in (wk.get("entries") or []) if not e.get("archived") or e.get("user_name")]
        show_mc = any(e.get("mc_fees") is not None for e in entries)
        ws = wb.create_sheet(tab_name(we, used))
        heads = ["Name", "PR/OR", "Coach", "Stage", "Goal", "LW", "B/E", *DAYS, "Total", "Gold", "Silver",
                 "Round", "Days", "P/A", "BA fees"] + (["MC fees"] if show_mc else [])
        ws["A1"] = f"{office_name or org_name} · Bells · week ending {uk_date(we)}"
        ws["A1"].font = Font(bold=True, size=13)
        for c, h in enumerate(heads, start=1):
            cell = ws.cell(row=3, column=c, value=h)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=HEAD_FILL)
            cell.alignment = centre
            cell.border = box
        row = 4

        # Teams in the order the sheet meets them.
        teams: dict[str, list] = {}
        for e in entries:
            teams.setdefault(team_of(e, org_name), []).append(e)

        def totals(rows: list) -> dict:
            t = {"days": [0] * 7, "total": 0, "gold": 0, "silver": 0, "round": 0, "worked": 0,
                 "fees": 0.0, "mc": 0.0, "lw": None}
            for e in rows:
                for i, d in enumerate((e.get("days") or [])[:7]):
                    if (d or {}).get("status") == "in" and not _exited_on(e, we, i):
                        t["days"][i] += int(_num(d.get("over30")) + _num(d.get("under30")))
                t["total"] += int(_num(e.get("total_sales")))
                t["gold"] += int(_num(e.get("total_over30")))
                t["silver"] += int(_num(e.get("total_under30")))
                t["round"] += int(_num(e.get("total_memberships")))
                t["worked"] += int(_num(e.get("days_worked")))
                t["fees"] += _num(e.get("earnings"))
                t["mc"] += _num(e.get("mc_fees"))
                if e.get("last_week_total") is not None:
                    t["lw"] = (t["lw"] or 0) + int(_num(e.get("last_week_total")))
            return t

        def total_row(label: str, t: dict, goal=None):
            nonlocal row
            vals = [label, "", "", "", goal if goal is not None else "", t["lw"] if t["lw"] is not None else "", "",
                    *t["days"], t["total"], t["gold"], t["silver"], t["round"], t["worked"],
                    round(t["total"] / t["worked"], 1) if t["worked"] else "", round(t["fees"], 2)]
            if show_mc:
                vals.append(round(t["mc"], 2))
            for c, v in enumerate(vals, start=1):
                cell = ws.cell(row=row, column=c, value=v)
                cell.font = Font(bold=True)
                cell.fill = PatternFill("solid", fgColor=TOTAL_FILL)
                cell.border = box
                if c > 1:
                    cell.alignment = centre
            row += 1

        office_rows: list = []
        for team, rows in teams.items():
            leader = next((e for e in rows if e.get("user_id") and e.get("user_id") == (e.get("primary_team") or {}).get("leader_id")), None)
            band = ws.cell(row=row, column=1, value=team.upper())
            band.font = Font(bold=True, color=HEAD_FILL)
            for c in range(1, len(heads) + 1):
                ws.cell(row=row, column=c).fill = PatternFill("solid", fgColor=TEAM_FILL)
            row += 1
            for e in rows:
                days = (e.get("days") or [])[:7]
                goal, lw = e.get("weekly_goal"), e.get("last_week_total")
                worked = int(_num(e.get("days_worked")))
                vals = [e.get("user_name") or "", e.get("recruit_badge") or "", e.get("coach_name") or "",
                        e.get("stage_label") or "", goal if goal is not None else "", lw if lw is not None else "",
                        e.get("break_even") if e.get("break_even") is not None else ""]
                for c, v in enumerate(vals, start=1):
                    cell = ws.cell(row=row, column=c, value=v)
                    cell.border = box
                    if c > 1:
                        cell.alignment = centre
                if goal is not None:
                    paint(ws.cell(row=row, column=5), WEEK_LOOK[week_band(_num(goal))])
                if lw is not None:
                    paint(ws.cell(row=row, column=6), WEEK_LOOK[week_band(_num(lw))])
                for i in range(7):
                    v, look = day_cell(e, days[i] if i < len(days) else {}, we, i)
                    cell = ws.cell(row=row, column=8 + i, value=v)
                    cell.alignment = centre
                    cell.border = box
                    paint(cell, look, bold=bool(look))
                rest = [int(_num(e.get("total_sales"))), int(_num(e.get("total_over30"))), int(_num(e.get("total_under30"))),
                        int(_num(e.get("total_memberships"))), worked,
                        round(_num(e.get("total_sales")) / worked, 1) if worked else "", round(_num(e.get("earnings")), 2)]
                if show_mc:
                    rest.append(round(_num(e.get("mc_fees")), 2))
                for k, v in enumerate(rest):
                    cell = ws.cell(row=row, column=15 + k, value=v)
                    cell.alignment = centre
                    cell.border = box
                    if k == 0:
                        cell.font = Font(bold=True)
                row += 1
            team_goal = (leader or {}).get("team_weekly_goal")
            if team == org_name and team_goal is None:
                team_goal = wk.get("mini_team_goal")
            total_row(f"{team} total", totals(rows), team_goal)
            office_rows += rows
            row += 1
        if entries:
            total_row(f"{office_name or org_name} total", totals(office_rows))

        # Money columns as £, widths, frozen header and names.
        money_cols = [len(heads)] + ([len(heads) - 1] if show_mc else [])
        for r in range(4, row):
            for c in money_cols:
                cell = ws.cell(row=r, column=c)
                if isinstance(cell.value, (int, float)):
                    cell.number_format = '"£"#,##0.00'
        widths = {1: 24, 2: 7, 3: 14, 4: 8}
        for c in range(1, len(heads) + 1):
            ws.column_dimensions[get_column_letter(c)].width = widths.get(c, 9 if c < len(heads) - 1 else 11)
        ws.freeze_panes = "B4"

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
